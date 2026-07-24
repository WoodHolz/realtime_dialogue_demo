# -*- coding: utf-8 -*-
"""tmp-prod Excel helpers (openpyxl)."""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

HEADER_ALIASES = {
    "用人单位": "employer",
    "姓名": "name",
    "性别": "gender",
    "手机号": "phone",
}


def _norm_header(cell: Any) -> str:
    return re.sub(r"\s+", "", str(cell or "").strip())


def _norm_phone(raw: Any) -> str:
    s = str(raw or "").strip()
    if s.endswith(".0") and s.replace(".", "", 1).isdigit():
        s = s[:-2]
    s = re.sub(r"[\s\-()]", "", s)
    if s.startswith("+86"):
        s = s[3:]
    return s


def read_contacts_from_xlsx(data: bytes | Path) -> list[dict[str, str]]:
    """Parse docs/tmp.xlsx-like sheets: find header row with 用人单位/姓名/手机号."""
    if isinstance(data, Path):
        wb = load_workbook(data, read_only=True, data_only=True)
    else:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    header_idx = -1
    colmap: dict[str, int] = {}
    for i, row in enumerate(rows):
        cells = [_norm_header(c) for c in (row or ())]
        found: dict[str, int] = {}
        for j, h in enumerate(cells):
            if h in HEADER_ALIASES:
                found[HEADER_ALIASES[h]] = j
        if {"employer", "name", "phone"}.issubset(found.keys()):
            header_idx = i
            colmap = found
            # gender optional
            if "gender" not in colmap and "性别" in cells:
                colmap["gender"] = cells.index("性别")
            break
    if header_idx < 0:
        raise ValueError("未找到表头（需要列：用人单位、姓名、手机号）")

    contacts: list[dict[str, str]] = []
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        def cell(key: str) -> str:
            idx = colmap.get(key)
            if idx is None or idx >= len(row):
                return ""
            return str(row[idx] or "").strip()

        name = cell("name")
        phone = _norm_phone(cell("phone") if "phone" in colmap else "")
        employer = cell("employer")
        gender = cell("gender")
        if not name and not phone:
            continue
        if not phone or not phone.isdigit() or len(phone) < 7:
            continue
        contacts.append(
            {
                "employer": employer,
                "name": name,
                "gender": gender,
                "phone": phone,
            }
        )
    return contacts


def call_task_to_result_row(task: Any) -> dict[str, Any]:
    """Map outbound CallTask → write_results_xlsx row (tmp-prod single / batch shape)."""
    payload = getattr(task, "payload", None) or {}
    if not isinstance(payload, dict):
        payload = {}
    telephony = getattr(task, "telephony", None) or {}
    if not isinstance(telephony, dict):
        telephony = {}
    sip = telephony.get("sip") or {}
    if not isinstance(sip, dict):
        sip = {}
    business = telephony.get("business") or {}
    if not isinstance(business, dict):
        business = {}
    recordings = sip.get("recordings") or {}
    if not isinstance(recordings, dict):
        recordings = {}
    stereo = str(recordings.get("stereo") or "")
    recording_dir = ""
    if stereo:
        try:
            recording_dir = str(Path(stereo).parent)
        except Exception:
            recording_dir = ""
    name = str(
        payload.get("candidate_name")
        or payload.get("name")
        or ""
    ).strip()
    return {
        "employer": str(payload.get("employer") or payload.get("client_name") or "").strip(),
        "name": name,
        "gender": str(payload.get("gender") or "").strip(),
        "phone": str(getattr(task, "phone", "") or payload.get("phone") or "").strip(),
        "status": str(getattr(task, "status", "") or ""),
        "result": str(getattr(task, "result", None) or business.get("result") or ""),
        "result_summary": str(
            getattr(task, "result_summary", "") or business.get("result_summary") or ""
        ),
        "attempt": 1,
        "fail_reason": str(getattr(task, "error", "") or ""),
        "promised_time": str(business.get("promised_time") or ""),
        "email": str(business.get("email") or ""),
        "last_call_id": str(getattr(task, "call_id", "") or ""),
        "transcript": list(getattr(task, "transcript", None) or []),
        "stereo_path": stereo,
        "recordings": recordings,
        "recording_dir": recording_dir,
        "updated_at": str(
            getattr(task, "finished_at", None)
            or getattr(task, "updated_at", "")
            or ""
        ),
    }


def write_results_xlsx(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "外呼结果"
    # 交付跟进 / 交付跟进结果放最后：前者系统回写「是」；后者留给交付手填（如是否再呼）。
    headers = [
        "用人单位",
        "姓名",
        "性别",
        "手机号",
        "状态",
        "结果",
        "结果说明",
        "尝试次数",
        "失败原因",
        "承诺完成时间",
        "登记邮箱",
        "最近call_id",
        "转写摘要",
        "双声道录音(L用户/R AI)",
        "录音目录",
        "更新时间",
        "交付跟进",
        "交付跟进结果",
    ]
    ws.append(headers)
    for r in rows:
        transcript = r.get("transcript") or []
        summary = ""
        if isinstance(transcript, list) and transcript:
            bits = []
            for t in transcript[-8:]:
                role = t.get("role", "")
                text = t.get("text", "")
                bits.append(f"{role}:{text}")
            summary = " | ".join(bits)[:2000]
        stereo = r.get("stereo_path") or ""
        if not stereo:
            recs = r.get("recordings") or {}
            if isinstance(recs, dict):
                stereo = str(recs.get("stereo") or "")
        result = str(r.get("result") or "")
        # 仅拒提供邮箱 → 标记需交付跟进；其余为空。
        delivery_followup = "是" if result == "no_email_feedback" else ""
        # 交付跟进结果：留给交付手填（是否再呼等），系统不写。
        delivery_followup_result = str(r.get("delivery_followup_result") or "")
        ws.append(
            [
                r.get("employer", ""),
                r.get("name", ""),
                r.get("gender", ""),
                r.get("phone", ""),
                r.get("status", ""),
                result,
                r.get("result_summary", ""),
                r.get("attempt", 0),
                r.get("fail_reason", ""),
                r.get("promised_time", ""),
                r.get("email", ""),
                r.get("last_call_id", ""),
                summary,
                stereo,
                r.get("recording_dir", ""),
                r.get("updated_at", ""),
                delivery_followup,
                delivery_followup_result,
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
